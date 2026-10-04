"""Create the formal report, truthful prompt log and recording guide.
Body formatting follows the supplied WAI policy: Times New Roman12,
1.5 spacing, justified text and one-inch margins. Explicit page breaks
control the report's twelve-page budget; render QA verifies actual pagination.
"""
from pathlib import Path
import json, re
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]; OUT=ROOT/'submission';OUT.mkdir(exist_ok=True)
S=json.loads((ROOT/'data/clean/source_register.json').read_text())
def doc(title):
    d=Document();sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
    sec.header_distance=sec.footer_distance=Inches(.45)
    for n in ['Normal','Body Text','Caption','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        st=d.styles[n];st.font.name='Times New Roman';st.font.size=Pt(12);st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.line_spacing=1.5;st.paragraph_format.space_after=Pt(3)
        st.paragraph_format.widow_control=True
        rf=st.element.get_or_add_rPr().get_or_add_rFonts()
        for key in list(rf.attrib):
            if 'Theme' in key: del rf.attrib[key]
        for key in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+key),'Times New Roman')
    d.styles['Normal'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    for n in ['Heading 1','Heading 2']:
        d.styles[n].font.bold=True;d.styles[n].paragraph_format.space_before=Pt(7);d.styles[n].paragraph_format.space_after=Pt(5)
    d.styles['Title'].font.size=Pt(18);d.styles['Title'].font.bold=True
    hp=sec.header.paragraphs[0];hp.text='IIM Ranchi  |  Operations Management  |  Working with AI';hp.style='Normal';hp.alignment=WD_ALIGN_PARAGRAPH.LEFT
    for r in hp.runs:r.font.size=Pt(9)
    fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=fp.add_run('Page ');r.font.name='Times New Roman';r.font.size=Pt(10)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');fp._p.append(fld)
    for st in d.styles:
        for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    d.core_properties.title=title;d.core_properties.subject='Operations Management WAI assignment';d.core_properties.author='Student author'
    return d
def p(d,t,boldlead=None):
    x=d.add_paragraph(t);x.paragraph_format.keep_with_next=False;return x
def h(d,t):return d.add_heading(t,1)
def page(d,t):
    x=h(d,t);x.paragraph_format.page_break_before=True
def pic(d,file,width=6.4):
    x=d.add_paragraph();x.alignment=WD_ALIGN_PARAGRAPH.CENTER;x.paragraph_format.space_after=Pt(2);x.add_run().add_picture(str(ROOT/'figures'/file),width=Inches(width));return x
def caption(d,t):
    x=d.add_paragraph(t,'Caption');x.paragraph_format.line_spacing=1.15;x.paragraph_format.space_after=Pt(6)
    for r in x.runs:r.italic=True;r.font.size=Pt(10)
def table(d,heads,rows,widths=None):
    t=d.add_table(rows=1,cols=len(heads));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if widths:
        for col,w in zip(t.columns,widths):col.width=Inches(w)
    for j,v in enumerate(heads):t.rows[0].cells[j].text=str(v)
    for row in rows:
        for cell,v in zip(t.add_row().cells,row):cell.text=str(v)
    for i,row in enumerate(t.rows):
        for j,c in enumerate(row.cells):
            if widths:c.width=Inches(widths[j])
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');b.append(el)
            pr.append(b);m=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'70');el.set(qn('w:type'),'dxa');m.append(el)
            pr.append(m)
            if i==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E9EEF0');pr.append(sh)
            for pp in c.paragraphs:
                pp.paragraph_format.line_spacing=1.5;pp.paragraph_format.space_after=Pt(0)
                pp.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                for rr in pp.runs:rr.font.size=Pt(12);rr.font.name='Times New Roman';rr.bold=i==0
        trpr=row._tr.get_or_add_trPr();ns=OxmlElement('w:cantSplit');trpr.append(ns)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    return t
def hyperlink(p,text,url):
    part=p.part;rid=part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    e=OxmlElement('w:hyperlink');e.set(qn('r:id'),rid);r=OxmlElement('w:r');pr=OxmlElement('w:rPr');fonts=OxmlElement('w:rFonts');fonts.set(qn('w:ascii'),'Times New Roman');pr.append(fonts)
    sz=OxmlElement('w:sz');sz.set(qn('w:val'),'24');pr.append(sz);r.append(pr);t=OxmlElement('w:t');t.text=text;r.append(t);e.append(r);p._p.append(e)
def ref(d,s):
    x=d.add_paragraph();x.paragraph_format.left_indent=Inches(.3);x.paragraph_format.first_line_indent=Inches(-.3)
    x.paragraph_format.space_after=Pt(9);x.alignment=WD_ALIGN_PARAGRAPH.LEFT
    year=s['year'].replace('n.d.-a','n.d.')
    x.add_run(s['author'].rstrip('.')+'. ('+year+'). ');r=x.add_run(s['title']+'. ');r.italic=True
    if s['id']=='S8':x.add_run('OTexts. ')
    if s['id']=='S5':x.add_run('Press Information Bureau. ')
    if s['id']=='S9':x.add_run('Retrieved October 4, 2026, from ')
    # URLs are live hyperlinks. Break opportunities preserve readable line lengths.
    visible=s['url'].replace('%20','%20\u200b').replace('/','/\u200b').replace('-','-\u200b').replace('_','_\u200b')
    hyperlink(x,visible,s['url'])

def build_report():
    d=doc('Dedicated freight corridors and operating performance')
    d.add_paragraph('Dedicated freight corridors and operating performance','Title')
    p(d,'An AI assisted assessment of Indian Railways freight capacity and coordination')
    p(d,'Operations Management WAI assignment  |  4 October 2026')
    p(d,'Student name: ____________________    Roll number: ____________________')
    h(d,'Executive summary')
    p(d,'Dedicated freight corridors can separate freight from passenger traffic, but higher traffic volume alone does not establish better service or nationwide congestion relief. This study examines Eastern and Western Dedicated Freight Corridor operations using public administrative data and develops a limited, testable managerial response.')
    p(d,'Reported DFC train movements increased from 88,225 in FY2023–24 to 130,116 in FY2024–25, a 47.5% rise. Gross tonne-kilometres increased by 61.3%. Over the same period, EDFC’s unweighted mean monthly speed fell from 45.11 to 41.65 km/h, while mean reported wagon turnround rose from 1.22 to 1.64 days. These indicators justify investigation of operational interfaces; they do not identify the cause of delay (Dedicated Freight Corridor Corporation of India Limited [DFCCIL], 2025).')
    p(d,'A day-adjusted naive forecast achieved 4.66% weighted absolute percentage error on twelve rolling monthly test forecasts, compared with 9.94% for random forest. Historical diversion scenarios show that relief varies substantially by conventional section. They cannot establish how much capacity was actually released today.')
    p(d,'The recommendation is a twelve-week EDFC interchange coordination pilot, supported by timestamp capture, shared KPIs and human review. Under explicitly hypothetical assumptions, its first-year net benefit is INR 0.50 crore. Approval should depend on measured delays and validated costs, not this illustrative return.')
    p(d,'The accompanying package contains audited data, commented Python code, seven dashboard visuals, an evidence-retrieval chatbot, a prompt logbook and ten slides. All observed-data calculations are reproducible; no synthetic observations were added.')

    page(d,'1  Problem and managerial relevance')
    p(d,'Indian Railways and the Dedicated Freight Corridor Corporation of India Limited (DFCCIL) must convert dedicated capacity into reliable freight flow. This requires coordination with conventional routes, terminals and rolling-stock operations.')
    p(d,'The study addresses three questions. How have DFC traffic, speed and wagon turnround evolved? Does a small machine-learning model improve short-horizon traffic forecasts? What does published conventional-line traffic imply about possible capacity relief under stated assumptions?')
    h(d,'Operations management framework')
    p(d,'Capacity management distinguishes infrastructure from effective throughput. Bottleneck analysis asks whether an interchange, terminal, crew handover or downstream section limits the journey. The theory of constraints motivates a targeted pilot; it does not prove which interface caused the observed changes.')
    p(d,'Little’s Law links work in process, throughput and flow time when boundaries and units match. Published train counts and wagon turnround do not meet that requirement, so they cannot establish fleet needs. Aggregate monthly data also cannot calibrate a credible train-dispatch queueing model.')
    h(d,'Scope and stakeholders')
    p(d,'Core corridor comparisons cover April 2021–March 2025; total train movements extend back to December 2020. Conventional capacity scenarios use ten FY2022–23 sections reported by the World Bank. Freight customers, terminal operators, railway controllers, crews, infrastructure managers and passengers have interests in reliability, capacity and safe maintenance access.')
    p(d,'The managerial value is better prioritization of investigation and pilot measurement. National causal attribution, live scheduling and whole-project financial appraisal are outside the evidence.')
    h(d,'Final problem statement')
    p(d,'DFC traffic is growing, while EDFC speed and wagon turnround show operational pressure. This project uses audited public data, forecasting and transparent scenarios to identify a defensible coordination pilot and the evidence required to evaluate it.')

    page(d,'2  Literature and evidence review')
    p(d,'Ten sources cover operational records, evaluations, logistics research and technical methods. Measured observations, managerial arguments and appraisal targets are kept distinct.')
    p(d,'DFCCIL’s latest annual report supplies monthly operations data; the previous report enables vintage checks and describes its freight information system (DFCCIL, 2024, 2025). A March 2023 speed conflict demonstrates the need for extraction audits. Both reports originate from the same organization.')
    p(d,'The World Bank’s completion report provides section-level conventional traffic and utilization, enabling bounded diversion arithmetic (World Bank, 2024). The Indian Railways yearbook supplies a national goods-speed benchmark, but its traffic population and aggregation differ from corridor monthly averages (Ministry of Railways, n.d.). The Ministry’s year-end review independently publishes October 2024 DFC figures that agree with the annual-report cells; publication agreement is not independent field measurement (Ministry of Railways, 2024).')
    p(d,'BRIEF’s NITI-funded study identifies operational, terminal and connectivity constraints; NITI leaves responsibility for findings with BRIEF (Bureau of Research on Industry and Economic Fundamentals, n.d.). The World Bank’s rail-logistics announcement also emphasizes multimodal terminals and connectivity (World Bank, 2022). Together they support investigation beyond mainline capacity.')
    p(d,'JICA’s ex-ante evaluation explains the rationale for dedicated freight capacity and locomotive investment. Its future targets are planning assumptions, not evidence that benefits were achieved (Japan International Cooperation Agency, 2017). This distinction guides the treatment of the present project’s cost and diversion scenarios.')
    p(d,'Hyndman and Athanasopoulos (2021) establish why forecast evaluation must preserve time order. The scikit-learn documentation explains the random-forest ensemble used as an ML comparator (scikit-learn developers, n.d.). These methods justify an experiment, not an assumption that ML will outperform a simpler model.')
    p(d,'The gap is a matched panel of journeys, terminal dwell and conventional-route outcomes around commissioning. Public data supports monitoring and conditional scenarios. This project connects those observations to model evaluation and a testable intervention.')

    page(d,'3  Methodology and data preparation')
    pic(d,'06_methodology.png',5.6);caption(d,'Figure 1. Workflow from source collection to reviewed managerial recommendations.')
    p(d,'ChatGPT supported source discovery, code drafting, interpretation and visualization design. Python extracted tables, audited arithmetic and produced deterministic outputs. Pandas handled the monthly panel; scikit-learn supplied random forest and TF-IDF; Matplotlib generated figures. An offline HTML dashboard presents results and a retrieval chatbot. Source-first prompts progressed from feasibility to implementation; the logbook separates actual prompts from suggested follow-ups.')
    p(d,'The workflow extracted 444 reported metric cells and 52 monthly train counts. The clean corridor panel contains 96 corridor-month rows, with speed, gross tonne-kilometres and wagon turnround for four years. Corridor trains/day begins in April 2022 in the selected table. Ten conventional-section records were manually transcribed and checked against the World Bank table.')
    p(d,'Financial years run April–March. Dates, units and source locators are retained. Annual network trains/day equals annual counts divided by calendar days, including leap years. Corridor annual speed and turnround are unweighted means of reported monthly averages; they are not official distance- or time-weighted annual measures.')
    p(d,'Missing corridor rates and disputed cells remain missing. No forward fill, zero fill or synthetic observation is used. Tree models require no numerical standardization. TF-IDF normalizes evidence-text vectors for retrieval. Scenarios use normalized utilization, retaining their assumptions separately from observed data.')

    page(d,'4  Traffic growth and data quality')
    pic(d,'01_train_growth.png');caption(d,'Figure 2. DFC monthly train movements. Source: DFCCIL (2025), p. 48, Table 3.')
    table(d,['Metric','FY2023–24','FY2024–25'],[['Train movements','88,225','130,116'],['Trains per calendar day','241.05','356.48'],['GTKM in millions','119,128','192,144']],[3.1,1.7,1.7])
    p(d,'Traffic rose 47.5% and GTKM rose 61.3%. Gross tonne-kilometres combine train weight and distance; they are not net cargo tonnes. Expansion of the operating network means neither increase can be read as a same-network productivity gain. FY2020–21 is a partial December–March period and is excluded from annual growth comparisons.')
    p(d,'All five annual train totals reconcile to their monthly counts. Three material corrections to analytical use were necessary. February 2023 corridor rates sum to 222.00 trains/day, against a reported total of 229.93, so both corridor cells are excluded. The FY2023–24 Table 2 total average repeats April’s 211.67, so annual rates are recomputed. WDFC March 2023 speed differs across report vintages, so that cell is excluded.')
    p(d,'GTKM cells also exhibit small rounding inconsistencies. Network totals use the published monthly total column; corridor totals use their own monthly columns. They are not forced to agree. October 2024 count, GTKM and speeds match the PIB table, providing a useful cross-publication check (Ministry of Railways, 2024).')

    page(d,'5  Speed and wagon turnround')
    pic(d,'02_speed_wtr.png');caption(d,'Figure 3. Monthly reported speed and wagon turnround. Source: DFCCIL (2025), p. 49.')
    table(d,['Corridor and measure','FY2023–24','FY2024–25'],[['EDFC speed in km/h','45.11','41.65'],['WDFC speed in km/h','54.11','51.97'],['EDFC turnround in days','1.22','1.64'],['WDFC turnround in days','0.80','0.78']],[3.1,1.7,1.7])
    p(d,'These are unweighted means of twelve monthly values. EDFC turnround increased 34.2%, while its mean speed fell 7.7%. The implication is to investigate where time is spent as traffic grows. Longer journeys, commodity mix, new sections, maintenance and interchange delays are possible explanations; the available data cannot separate them.')
    p(d,'EDFC traffic and speed have a level correlation of −0.86 across 35 valid months. The correlation falls to −0.31 across 33 valid monthly changes. That reduction warns against interpreting a shared time trend as a causal relationship. WDFC’s level correlation is approximately 0.04.')
    p(d,'The national broad-gauge goods-speed benchmark is 25.0 km/h in FY2023–24 (Ministry of Railways, n.d.). It describes a different operating population. The study therefore reports it as context without claiming a matched DFC treatment effect. Excluding the disputed WDFC March 2023 value yields a FY2022–23 mean of 51.59 km/h; using the old or new value gives 51.35 or 52.18 respectively.')

    page(d,'6  Forecasting and AI evaluation')
    p(d,'Five models predict total monthly DFC train movements one month ahead. Validation covers April 2023–March 2024; the held-out test covers April 2024–March 2025. Each origin uses only previous observations, while earlier observed test months can update later one-step forecasts. This is retrospective data evaluation; historical publication availability was not reconstructed (Hyndman & Athanasopoulos, 2021).')
    table(d,['Model','Validation MAE','Test MAE','Test WAPE'],[['Naive','395','804','7.42%'],['Day-adjusted naive','298','505','4.66%'],['Seasonal naive','2,162','3,491','32.20%'],['12-month drift','377','846','7.80%'],['Random forest','586','1,078','9.94%']],[2.45,1.35,1.2,1.5])
    caption(d,'Table 3. MAE is mean absolute error in train movements. WAPE is total absolute error divided by total actual movements. Twelve predictions per split.')
    p(d,'The random forest uses three lags, a prior-three-month mean, time index, month sine/cosine and known month length. It has 200 trees, maximum depth 3, minimum leaf size 3 and seed 42, fixed before evaluation. Small samples and difficulty extrapolating growth limit its usefulness here (scikit-learn developers, n.d.).')
    p(d,'The validation winner is the day-adjusted naive model: last month’s movements per calendar day multiplied by next month’s days. Its test MAE is 505 trains and RMSE is 581. The April 2025 demonstration forecast is 12,094, using data through March 2025. It is not a live 2026 forecast. An experimental validation-residual band covered only 5 of 12 test outcomes; it is unsuitable for operational commitments.')
    p(d,'TF-IDF extracts distinctive terms from eleven curated source summaries, organized into themes such as connectivity, capacity and data quality. The dashboard chatbot retrieves these cited cards and abstains on weak matches. Neither theme frequency nor retrieval relevance establishes causal evidence. Managerially, simple monitoring is justified; autonomous dispatch is not.')

    page(d,'7  Conventional capacity relief scenarios')
    pic(d,'04_capacity.png',5.6);caption(d,'Figure 4. Conditional diversion arithmetic using FY2022–23 traffic. Source: World Bank (2024), Annex 4, Table 5.')
    p(d,'The reported conventional-line data already includes some DFC operation. It is therefore not a clean pre-DFC baseline. The scenario holds passenger and other services fixed and diverts an additional fraction of the remaining historical freight. It assumes equal train weights and unchanged effective capacity.')
    p(d,'Scenario utilization equals reported utilization multiplied by remaining total trains divided by original total trains. The denominator is normalized to the published utilization, because reported nominal capacity and the maintenance-adjusted utilization percentage do not give interchangeable arithmetic. The model does not estimate engineering headways or feasible timetable paths.')
    table(d,['Section','Reported','50% diversion','100% diversion'],[['DDU–Allahabad','148.0%','119.8%','91.5%'],['Kanpur–Tundla','129.0%','118.0%','107.0%'],['Khurja–Dadri','126.0%','109.6%','93.2%']],[2.6,1.1,1.5,1.3])
    p(d,'In this mechanical scenario, DDU–Allahabad needs roughly 85% additional freight diversion to reach a 100% index. Kanpur–Tundla remains above 100% even after all reported freight is removed. This does not describe current conditions; it illustrates why diversion alone may not remove every bottleneck. Section train counts must not be summed as unique journeys.')

    page(d,'8  Recommendations and pilot economics')
    p(d,'Prioritize an EDFC interchange coordination pilot. First capture arrival, acceptance, release and departure timestamps, coded delay causes and route/commodity identifiers. Then align acceptance windows, crew readiness and terminal release with a daily joint exception review. Retain controller authority and protect maintenance blocks. Expand only if measured performance improves without displacing delay downstream.')
    p(d,'The following cost–benefit model is entirely illustrative. Assume 20 pilot trains/day, 300 operating days/year and INR 5,000 value per train-hour. Setup costs are INR 0.40 crore and annual operating costs INR 0.30 crore. Valued time savings are multiplied by an assumed cash-realization fraction; no freight revenue or DFC construction benefit is added.')
    table(d,['Assumed case','Hours saved','Cash share','Year 1 net'],[['Low','0.25','20%','−0.55 cr'],['Base','1.00','40%','0.50 cr'],['High','1.50','60%','2.00 cr']],[2.0,1.4,1.4,1.7])
    pic(d,'05_cost_sensitivity.png',5.8);caption(d,'Figure 5. First-year sensitivity with all other hypothetical inputs held fixed. INR crore.')
    p(d,'Base-case gross time value is INR 3.00 crore; realized benefit is INR 1.20 crore and annual net benefit is INR 0.90 crore. Simple payback is 5.3 months after steady-state benefits begin. Year-one break-even requires 35 minutes saved per train at 40% cash realization. Ramp-up, tax, financing and discounting are excluded; the base case is not a budget commitment.')
    p(d,'Longer-term terminal and junction investment should follow measured bottlenecks. BRIEF and World Bank (2022) support this connectivity focus. Operations and finance must verify that released time produces actual benefits.')

    page(d,'9  Implementation and responsible AI')
    pic(d,'07_implementation.png',5.9);caption(d,'Figure 6. Proposed twelve-week pilot after managerial approval, not the academic submission schedule.')
    p(d,'Assign an operations sponsor, analyst, data engineer and interchange/terminal supervisors. Train staff on timestamp definitions, delay coding, interpretation and override responsibilities. The assumed INR 0.40 crore setup budget allocates 0.20 to integration, 0.08 to analytics, 0.06 to training and 0.06 to contingency.')
    p(d,'Weeks 1–4 establish definitions and a baseline. Weeks 5–9 test coordination, with matched route, commodity and time-of-day comparisons. Weeks 8–10 check whether delays move to adjacent sections. Weeks 11–12 review the evidence. Proposed gates are a 10% reduction in median interchange dwell, no deterioration in the 90th percentile or safety/maintenance indicators, and positive verified economics. These are management targets, not forecast results.')
    p(d,'DFCCIL, railway controllers and terminal partners must jointly own the pilot. Avoid ranking individual staff with the prototype. Address missing timestamps, inconsistent identifiers and extra data-entry burden through minimal fields, agreed definitions and weekly sample audits.')
    p(d,'AI hallucination is controlled through source locators, raw-cell preservation and executable calculations. Administrative reporting and selected-source bias remain. No personal or confidential operational records are used in this study. Any future internal data integration needs authorized access, minimized identifiers and an agreed retention policy. Public accessibility does not remove attribution or reuse obligations.')
    p(d,'AI assisted research, coding and drafting. The student must review the evidence and supply the recording and repository link. No completed interview, deployment or realized saving is claimed.')

    page(d,'References')
    ordered=[S[i] for i in [5,1,0,7,6,3,4,8,9,2]]
    for s in ordered[:5]:ref(d,s)
    page(d,'References continued and annexure')
    for s in ordered[5:]:ref(d,s)
    h(d,'Annexure A  Reproducibility and submission map')
    p(d,'Raw and clean data, the dictionary and discrepancy log are under data/. Python scripts are under code/; numeric results are under outputs/. Open dashboard/DFC_Operations_Dashboard.html offline. Source excerpts are under source_documents/. The separate prompt logbook records actual prompts; the recording guide supplies the three-minute script and code explanation. Read START_HERE.md before submitting the package and repository link.')
    path=OUT/'WAI_DFC_Report.docx';d.save(path);return path

def build_log():
    d=doc('AI prompt logbook and research decisions');d.add_paragraph('AI prompt logbook and research decisions','Title')
    p(d,'This log records the user’s actual instructions available in the conversation. It does not reconstruct unrecorded prompts or claim that the student personally performed automated checks. Reflections are prepared for student review. A separate workflow record explains the analytical implementation.')
    entries=[
    ('1  Assignment interpretation','I am working on my Operations Management WAI assignment and really wish to excel at it. I am attaching the exact instructions from my professor. Go through each and every detail and attention to detail should be apex. Firstly go through everything, give me a brief understanding of the task from your end and feel free to ask any questions, if any.','Mapped the professor’s policy to required evidence, report, dashboard, code, prompt log, slides, video and repository.','Used to build the compliance checklist and twelve-page report structure.','A broad ambition becomes actionable only after the requirements and evidence boundaries are explicit.'),
    ('2  Topic and research preferences',"1. I haven't yet, but I was thinking of analyzing IRCTC Freight trains operations vs now both IRCTC and DFC freight trains and their operations. Basically Indian Railways as a whole and how it helped decongest etc.\n2. Submission deadline - 04th October EOD\n3. No I want to use your help and expertise for web scrapping and finding all relevant information and Data\n4. Yes all but I want you to do all that in accordance to your needs. But I understand every tools and am comfortable",'Clarified the institutional scope as Indian Railways and DFCCIL and prioritized public operational data.','Set the topic, deadline and permission to collect data and use analytical tools.','The original IRCTC wording was corrected. The project must test what public data supports rather than assume decongestion.'),
    ('3  Feasibility first','first do a feasibility study that whether even sufficient data is available for study and is that data publicly and easily accessible or not','Downloaded official reports, tested public access and extracted sample monthly tables. Identified limits to nationwide causal analysis.','Selected a feasible descriptive and scenario-based design.','Availability is more than finding a webpage: useful data must be downloadable, extractable and comparable.'),
    ('4  Authorization to complete','Sure then let\'s proceed with all the process and analysis','Built extraction, audits, clean datasets, temporal forecast comparisons, capacity scenarios, economics, dashboard and submission materials.','Authorized full implementation without repeated approval pauses.','Model results must determine the recommendation. The simple forecast winning is a useful result, not a failure of the project.'),
    ('5  Resume saved work','You had completed audit and I think you were working on dashboard when the limit reached. Try to see if progress was saved and continue','Confirmed that raw/clean files, analysis outputs and the dashboard source were present; continued verification and writing.','Preserved continuity rather than restarting research or changing the analytical cutoff.','Saving raw sources, code and results makes interrupted work reproducible.')]
    for i,(title,prompt,summary,use,reflection) in enumerate(entries):
        if i:d.add_page_break()
        h(d,title);p(d,'Tool: ChatGPT with web retrieval and a Python/JavaScript execution environment.')
        h(d,'Prompt text')
        for line in prompt.split('\n'): p(d,line)
        h(d,'Output summary');p(d,summary);h(d,'Usage in project');p(d,use);h(d,'Reflection');p(d,reflection)
    page(d,'6  Implementation record and critical decisions')
    p(d,'The entries below are a workflow record of assistant-executed analysis, not additional student prompts. Exact executable logic and output values are preserved in the code and outputs folders.')
    for title,text in [
      ('Source audit','The latest DFCCIL tables were parsed from pdftotext output. Annual train totals were reconciled, October 2024 was checked against PIB, and historical speed cells were compared with the prior annual report.'),
      ('Cleaning decision','Conflicting February 2023 corridor train rates and WDFC March 2023 speed were left missing. GTKM rounding differences and partial-year coverage were preserved in the discrepancy log.'),
      ('Forecast experiment','Five models used the same validation and test months. Validation MAE selected the day-adjusted naive model. Random forest did not outperform that baseline; the test year was not used to choose the model.'),
      ('Capacity and economics','Historical additional-diversion scenarios and pilot economics were separated from observed data. All monetary inputs were labeled assumptions. No national causal impact or realized return was asserted.'),
      ('Interpretation and communication','Figures, dashboard and narrative draw from the same numeric outputs. The chatbot returns cited summaries and can abstain. The report recommends measurement before scale-up.')]:
        h(d,title);p(d,text)
    page(d,'7  Suggested prompts for student review')
    p(d,'These are proposed follow-up prompts, not interactions already performed. Use them only after reviewing the files and add the actual outputs to your own log if you run them.')
    for t in [
      'Explain the difference between gross tonne-kilometres and cargo tonnes using the dataset, and identify claims that would confuse them.',
      'Audit whether each forecast feature was available before its prediction month. Show where the code prevents future information from entering the model.',
      'Challenge the interpretation that lower EDFC speed proves congestion. What additional route, distance and commodity data would be required?',
      'Recalculate first-year pilot economics with my own cost assumptions and tell me which inputs dominate the decision.',
      'Act as an Operations Management examiner. Ask one question at a time about capacity, bottlenecks, model validation and ethical AI, then assess my answer.'
    ]:p(d,t)
    p(d,'Student review note: add any further prompts you actually use, what you changed after reviewing the output, and your own reflection. Do not describe this prepared log as a complete platform export.')
    path=OUT/'WAI_DFC_Prompt_Logbook.docx';d.save(path);return path

SCRIPT=[
('0:00–0:20','Introduce yourself and show slide 1','Hello, I am [name], [roll number]. My Operations Management project examines whether dedicated freight corridor growth translates into better operating performance, and what public evidence can tell us about capacity relief on conventional Indian Railways lines.'),
('0:20–0:45','Show report source table, then dashboard charts 1–4','I used official DFCCIL reports, the Indian Railways yearbook and a World Bank evaluation. The dataset contains fifty-two monthly train counts and ninety-six corridor-month records. I preserved raw values and excluded conflicting cells. DFC train movements rose forty-seven point five percent, but EDFC speed fell and wagon turnround increased.'),
('0:45–1:10','Change dashboard year and corridor filters','Here, I can switch between corridors and years. The speed and turnround charts show why more traffic is not sufficient evidence of better service. These patterns justify investigating interchange coordination, but route length, commodity mix and phased commissioning also changed. The data does not prove that congestion caused the deterioration.'),
('1:10–1:40','Show run_analysis.py forecast function, then chart 5','In the Python code, each forecast uses only earlier observations. Three lagged counts and known calendar information feed the random forest. Model selection uses a validation year, followed by a separate test year. A simple day-adjusted naive forecast performed better: four point six six percent test error, compared with nine point nine four percent for random forest. Its calculation is last month’s daily rate multiplied by next month’s days.'),
('1:40–2:05','Move diversion slider from 0 to 50 and show chart 6','This slider explores additional diversion of historical freight. On DDU to Allahabad, the normalized utilization index falls from one hundred forty-eight percent to approximately one hundred twenty percent at fifty percent diversion. These are conditional calculations using a partly post-DFC baseline, not measured current capacity savings.'),
('2:05–2:30','Show cost sliders, then slide 9','I recommend a twelve-week interchange coordination pilot, starting with timestamp capture and matched movement comparisons. The illustrative base case gives fifty lakh rupees in first-year net benefit. However, every cost input is an assumption. The pilot must validate both time savings and their cash value before expansion.'),
('2:30–3:00','Ask the evidence assistant about capacity; finish on slide 10','The evidence assistant retrieves cited source summaries and declines weak matches. It is not a live dispatcher. ChatGPT helped with research, code and drafting, while explicit audits make the calculations reviewable. The managerial lesson is to combine capacity investment with reliable interfaces, and choose the simplest model that performs well. The key next step is measuring actual interchange delays before making a larger commitment.')]
def build_guide():
    d=doc('Presentation recording and viva guide');d.add_paragraph('Presentation recording and viva guide','Title')
    p(d,'Record one video of approximately three minutes. Replace the name and roll-number fields before recording. The script is a speaking draft, not proof that a recording or repository upload has already been completed.')
    p(d,'Open the slide deck, dashboard and code before starting. Record the screen with microphone audio. Use a small face-camera overlay if required by your instructor. Rehearse once with a timer; shorten pauses or phrasing to meet three minutes without making the code demonstration unreadable.')
    for timing,visual,spoken in SCRIPT:
        h(d,timing+'  '+visual);p(d,spoken)
    page(d,'Code explanation and demonstration steps')
    for title,text in [
      ('Extraction','extract_data.py locates report table headings, parses monthly values and attaches source IDs. It checks that all five annual totals equal the sums of monthly train counts.'),
      ('Missing values','The raw file preserves original numbers. The clean panel leaves disputed values missing; it never inserts zeros or invented speeds. This prevents apparent precision from concealing source uncertainty.'),
      ('Feature construction','features() builds three lagged counts, a three-month mean, a trend index, seasonal sine/cosine terms and the number of days in the target month. All lag values precede the prediction date.'),
      ('Temporal evaluation','forecast() slices y[:t]. The validation year selects the model using MAE; the test year reports later one-step performance. Earlier test observations can update subsequent forecasts, which is different from a fixed twelve-month forecast.'),
      ('Interpreting the winning model','Day-adjusted naive holds the most recent daily train rate constant. This handles unequal month lengths. It won this small retrospective experiment, but a larger, newer dataset could change the ranking.'),
      ('Dashboard demonstration','Open the HTML locally, select FY2024–25 and EDFC, change the forecast model, move the diversion slider to 50%, then show how first-year benefit changes when savings or cash realization fall.'),
      ('Text analysis and chatbot','Python TF-IDF ranks words in curated evidence cards. The browser uses token TF-IDF cosine similarity to retrieve cited cards. This is local retrieval, not an LLM generating unrestricted answers.')]:h(d,title);p(d,text)
    page(d,'Viva questions and concise answers')
    qa=[
      ('Why not compare IRCTC with DFC?','The study concerns Indian Railways freight operations and DFCCIL infrastructure. IRCTC is not the relevant freight network operator.'),
      ('Can you prove DFC caused national decongestion?','No. The data lacks a matched untreated panel and consistent pre/post section outcomes. The project supports descriptive findings and conditional capacity arithmetic.'),
      ('Why did speed fall while traffic rose?','The pattern is observed, but causes are not isolated. Journey length, traffic composition, commissioning and interface delays may all contribute.'),
      ('Why use AI if the simple model won?','AI supported research, code, text retrieval and testing. Critical evaluation includes rejecting an ML model when a simpler baseline performs better.'),
      ('Why not use random train/test splitting?','It would allow later traffic patterns to influence evaluation of earlier predictions. A chronological split better represents forecasting.'),
      ('What is WAPE?','The sum of absolute forecast errors divided by the sum of actual movements, expressed as a percentage. It is not classification accuracy.'),
      ('Is wagon turnround terminal dwell?','No. The public DFCCIL metric has its own operational boundary. The pilot proposes collecting dwell timestamps because that detail is missing.'),
      ('Why not use Little’s Law to estimate fleet size?','The published throughput and turnround series do not share clearly matched units and system boundaries. Multiplying them would be misleading.'),
      ('Are the pilot savings real?','They are hypothetical scenarios. Finance and operations must validate train-hour value, realized cash savings, costs and ramp-up.'),
      ('What is your strongest limitation?','Changing network coverage and sparse aggregate data prevent causal attribution and reliable operational uncertainty estimates.')]
    for q,a in qa:h(d,q);p(d,a)
    page(d,'Final submission checklist')
    for t in ['Add your name and roll number to the report and slide 1. Review every finding and be ready to explain the limitations.','Record the three-minute video with your own introduction, tool demonstration, code explanation and insights.','Upload the organized project folder to GitHub or Google Drive. Add the actual repository or folder URL to SUBMISSION_CHECKLIST.md and the LMS submission.','Check that the assessor can open the link with the intended access permissions. A local ZIP download is not the required GitHub/Drive link.','Submit the report, prompt logbook, commented code, dashboard, slide deck, video, data files and repository link by 4 October EOD in the required time zone.','Open the dashboard in a desktop browser. It is self-contained and needs no API key or internet connection for its calculations.','Do not claim that synthetic scenario inputs are measured DFCCIL costs, or that retrospective forecasts are live predictions.']:p(d,'☐ '+t)
    path=OUT/'WAI_DFC_Recording_and_Viva_Guide.docx';d.save(path);return path

if __name__=='__main__':
    for func in [build_report,build_log,build_guide]:print(func())
